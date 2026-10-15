"""참가 신청서 양식(docs/signup-template.xlsx) 생성.  python3 tools/make-signup-template.py
제45회 신청서와 같은 틀(A~M열, A4 세로 2장):
  1조·2조 가로 블록(각 20명), 3조(20명)는 1조 아래, 그 옆에 참가비 내역(수식)과 3인조(이름·성별·핸디·비고),
  아래(2쪽)에 스카치(왼쪽)·베이커(오른쪽) 각 10팀을 나란히, 팀마다 팀원 한 명씩 세로로(팀 | 이름 | 성별 | 핸디 | 비고).
핸디는 한 칸(여성·시니어·프로 감점 등 게임당 핸디 합계). 팀 핸디는 팀원 핸디의 합으로 적용된다.
수식·라벨 칸은 시트 보호로 잠그고 노란 입력 칸만 열어 둔다. js/signup.js 의 블록 파서가 그대로 읽는다."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.utils import get_column_letter as L

FONT = '맑은 고딕'
yellow = PatternFill('solid', fgColor='FFF2CC'); navy = PatternFill('solid', fgColor='1E2537'); grey = PatternFill('solid', fgColor='F2F2F2'); gold = PatternFill('solid', fgColor='FFFF00'); dim = PatternFill('solid', fgColor='E7E6E6')
thin = Side(style='thin', color='999999'); box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal='center', vertical='center'); left = Alignment(horizontal='left', vertical='center'); right = Alignment(horizontal='right', vertical='center')
UNLOCKED = Protection(locked=False)
GROUP_ROWS = 20; TEAMS = 10; LASTCOL = 13  # M
FEES = [('개인전 참가비', 35000), ('챔프전 참가비', 10000), ('사이드 참가비', 10000), ('클럽 참가비', 110000), ('클럽 찬조금', 50000)]

def f(size=9, bold=False, color='000000'): return Font(name=FONT, size=size, bold=bold, color=color)
def hcell(ws, r, c, v=None, fill=navy, color='FFFFFF', size=9):
    x = ws.cell(row=r, column=c)
    if v not in ('', None): x.value = v
    x.font = f(size, True, color); x.fill = fill; x.alignment = center; x.border = box; return x
def icell(ws, r, c, v=None, align=center, fmt=None):
    x = ws.cell(row=r, column=c)
    if v is not None: x.value = v
    x.font = f(); x.fill = yellow; x.border = box; x.alignment = align; x.protection = UNLOCKED
    if fmt: x.number_format = fmt
    return x
def bordered(ws, r, c1, c2):
    for c in range(c1, c2 + 1):
        x = ws.cell(row=r, column=c); x.border = box
        if x.font.name != FONT: x.font = f()
def merge(ws, r, c1, c2, r2=None): ws.merge_cells(start_row=r, start_column=c1, end_row=r2 or r, end_column=c2)
def memo(ws, r, c1, c2, text, r2=None, size=7.5):
    merge(ws, r, c1, c2, r2); x = ws.cell(row=r, column=c1, value=text); x.font = f(size, color='595959'); x.alignment = Alignment(wrap_text=True, vertical='top'); return x

def build(ws, example):
    ws.sheet_view.showGridLines = False
    widths = {'A': 6.5, 'B': 4.5, 'C': 10, 'D': 4.5, 'E': 6, 'F': 5.5, 'G': 5.5, 'H': 9, 'I': 10, 'J': 4.5, 'K': 7.5, 'L': 5.5, 'M': 9.5}
    for k, v in widths.items(): ws.column_dimensions[k].width = v
    for r in range(1, 90): ws.row_dimensions[r].height = 14
    merge(ws, 1, 1, LASTCOL); t = ws.cell(row=1, column=1, value='제46회 네이버 볼링 동호회 서울 전국대회 참가 신청서'); t.font = f(14, True); t.alignment = center; ws.row_dimensions[1].height = 24
    merge(ws, 2, 1, LASTCOL); n = ws.cell(row=2, column=1, value='노란 칸만 입력 · 3인조/팀 명단에는 개인전 명단과 똑같은 이름 · 챔프전은 1조만 · 사이드·챔프전 참가는 O 표시 · 수량·금액은 자동 계산')
    n.font = f(8, color='595959'); n.alignment = center
    merge(ws, 3, 1, 2); hcell(ws, 3, 1, '클럽명', grey, '000000'); merge(ws, 3, 3, 5); icell(ws, 3, 3, '드림존' if example else None, left, '@'); bordered(ws, 3, 4, 5)
    merge(ws, 3, 6, 7); hcell(ws, 3, 6, '지역명', grey, '000000'); merge(ws, 3, 8, 9); icell(ws, 3, 8, '청주' if example else None, left, '@'); bordered(ws, 3, 9, 9)
    merge(ws, 3, 10, 11); hcell(ws, 3, 10, '연락처', grey, '000000'); merge(ws, 3, 12, 13); icell(ws, 3, 12, '010-0000-0000 (총무 홍길동)' if example else None, left, '@'); bordered(ws, 3, 13, 13)
    # 개인전: B 순번 | C~H 1조 (성명 성별 핸디 사이드 챔프전 비고) | I~M 2조 (성명 성별 핸디 사이드 비고)
    first, last = 6, 6 + GROUP_ROWS - 1
    g3a, g3b = last + 2, last + 1 + GROUP_ROWS
    merge(ws, 4, 1, 1, g3b); a = hcell(ws, 4, 1, '개인전', grey, '000000', 10); a.alignment = Alignment(horizontal='center', vertical='center', text_rotation=255)
    hcell(ws, 4, 2, '순번', grey, '000000')
    merge(ws, 4, 3, 8); hcell(ws, 4, 3, '1조', grey, 'FF5A2E', 10); bordered(ws, 4, 4, 8)
    merge(ws, 4, 9, 13); hcell(ws, 4, 9, '2조', grey, 'FF5A2E', 10); bordered(ws, 4, 10, 13)
    h1 = ['성명', '성별', '핸디', '사이드', '챔프전', '비고']; h2 = ['성명', '성별', '핸디', '사이드', '비고']
    hcell(ws, 5, 2, None, grey)
    for i, h in enumerate(h1): hcell(ws, 5, 3 + i, h)
    for i, h in enumerate(h2): hcell(ws, 5, 9 + i, h)
    s1 = [('홍길동', '남', None, 'O', 'O', None), ('김영희', '여', 15, 'X', None, None), ('박철수', '남', 3, 'O', 'O', '시니어 63세')]
    s2 = [('이순자', '여', 20, 'O', '시니어 66세 (15+5)'), ('최민수', '남', -7, 'X', '프로')]
    s3 = [('정미라', '여', 15, 'O', None, None), ('강동원', '남', None, 'O', None, None)]
    def block_rows(r0, r1, c0, heads, sample, numcol=True):
        for r in range(r0, r1 + 1):
            if numcol: no = ws.cell(row=r, column=2, value=r - r0 + 1); no.font = f(9, color='808080'); no.border = box; no.alignment = center
            for i, h in enumerate(heads):
                is_text = h in ('성명', '비고'); icell(ws, r, c0 + i, None, left if is_text else center, '@' if is_text else None)
            if example and r - r0 < len(sample):
                for i, v in enumerate(sample[r - r0]):
                    if v is not None: ws.cell(row=r, column=c0 + i, value=v)
    block_rows(first, last, 3, h1, s1); block_rows(first, last, 9, h2, s2, numcol=False)
    merge(ws, last + 1, 3, 8); hcell(ws, last + 1, 3, '3조', grey, 'FF5A2E', 10); bordered(ws, last + 1, 4, 8); hcell(ws, last + 1, 2, None, grey)
    block_rows(g3a, g3b, 3, h1, s3)
    for r in range(g3a, g3b + 1): x = ws.cell(row=r, column=7); x.fill = dim; x.protection = Protection(locked=True)   # 3조에는 챔프전 없음
    def dv(formula, *ranges):
        d = DataValidation(type='list', formula1=formula, allow_blank=True)
        for rg in ranges: d.add(rg)
        ws.add_data_validation(d)
    dv('"남,여"', f'D{first}:D{last}', f'D{g3a}:D{g3b}', f'J{first}:J{last}')
    dv('"O,X"', f'F{first}:F{last}', f'F{g3a}:F{g3b}', f'L{first}:L{last}', f'G{first}:G{last}')
    # 참가비 내역 (2조 아래, 3조 옆): I:J 항목 | K 단가 | L 수량 | M 금액
    names = f'COUNTA(C{first}:C{last})+COUNTA(I{first}:I{last})+COUNTA(C{g3a}:C{g3b})'
    fr = last + 1
    merge(ws, fr, 9, 13); hcell(ws, fr, 9, '참가비 내역 (자동 계산)', grey, '000000', 10); bordered(ws, fr, 10, 13)
    merge(ws, fr + 1, 9, 10); hcell(ws, fr + 1, 9, '항목'); hcell(ws, fr + 1, 10); hcell(ws, fr + 1, 11, '단가'); hcell(ws, fr + 1, 12, '수량'); hcell(ws, fr + 1, 13, '금액')
    qty = {'개인전 참가비': f'={names}', '챔프전 참가비': f'=COUNTIF(G{first}:G{last},"O")', '사이드 참가비': f'=COUNTIF(F{first}:F{last},"O")+COUNTIF(L{first}:L{last},"O")+COUNTIF(F{g3a}:F{g3b},"O")', '클럽 참가비': 1, '클럽 찬조금': 1}
    for i, (label, unit) in enumerate(FEES):
        r = fr + 2 + i
        merge(ws, r, 9, 10); lab = ws.cell(row=r, column=9, value=label); lab.font = f(); lab.border = box; lab.alignment = left; ws.cell(row=r, column=10).border = box
        icell(ws, r, 11, unit, right, '#,##0')
        q = ws.cell(row=r, column=12, value=qty[label]); q.font = f(); q.border = box; q.alignment = center
        m = ws.cell(row=r, column=13, value=f'=K{r}*L{r}'); m.font = f(); m.border = box; m.alignment = right; m.number_format = '#,##0'
    r = fr + 2 + len(FEES)
    merge(ws, r, 9, 10); lab = ws.cell(row=r, column=9, value='개인 찬조'); lab.font = f(); lab.border = box; lab.alignment = left; bordered(ws, r, 10, 12)
    icell(ws, r, 13, 30000 if example else 0, right, '#,##0')
    tr = r + 1
    merge(ws, tr, 9, 12); hcell(ws, tr, 9, '총금액', gold, '000000', 10); bordered(ws, tr, 10, 12)
    tot = ws.cell(row=tr, column=13, value=f'=SUM(M{fr + 2}:M{r})'); tot.font = f(10, True); tot.fill = gold; tot.border = box; tot.alignment = right; tot.number_format = '#,##0'
    memo(ws, tr + 1, 9, 13, '단가는 대회 공지 기준으로 수정. 핸디는 게임당 합계(여성 15, 시니어 1~5, 프로 감점 등 합산)로 적고 내역은 비고에.', tr + 2)
    # 3인조 (참가비 내역 아래, 3조 옆): I 이름 | J 성별 | K 핸디 | L:M 비고
    rr = tr + 4
    hcell(ws, rr, 9, '3인조', navy, 'FFFFFF', 10); merge(ws, rr, 10, 13); x = ws.cell(row=rr, column=10, value='지역 대표 3명 (개인전 합산 → 지역 포인트)'); x.font = f(8, color='595959'); x.alignment = left; bordered(ws, rr, 10, 13)
    hcell(ws, rr + 1, 9, '이름'); hcell(ws, rr + 1, 10, '성별'); hcell(ws, rr + 1, 11, '핸디'); merge(ws, rr + 1, 12, 13); hcell(ws, rr + 1, 12, '비고'); hcell(ws, rr + 1, 13)
    reps = [('홍길동', '남', None, None), ('박철수', '남', 3, '시니어'), ('최민수', '남', -7, '프로')]
    for i in range(3):
        r = rr + 2 + i
        icell(ws, r, 9, None, left, '@'); icell(ws, r, 10); icell(ws, r, 11); merge(ws, r, 12, 13); icell(ws, r, 12, None, left, '@'); ws.cell(row=r, column=13).border = box
        if example:
            for c, v in zip((9, 10, 11, 12), reps[i]):
                if v is not None: ws.cell(row=r, column=c, value=v)
    dv('"남,여"', f'J{rr + 2}:J{rr + 4}')
    # 총 인원 요약
    sr = g3b + 1
    merge(ws, sr, 2, LASTCOL)
    summ = ws.cell(row=sr, column=2, value=f'="총 인원 "&({names})&"명 (1조 "&COUNTA(C{first}:C{last})&"명, 2조 "&COUNTA(I{first}:I{last})&"명, 3조 "&COUNTA(C{g3a}:C{g3b})&"명, 챔프전 "&L{fr + 3}&"명, 사이드 "&L{fr + 4}&"명)"')
    summ.font = f(9, True); summ.alignment = left; bordered(ws, sr, 2, LASTCOL)
    # 스카치(왼쪽 B~G) · 베이커(오른쪽 H~M): 팀 | 이름 | 성별 | 핸디 | 비고, 팀원 한 명씩 세로
    tl = sr + 1
    merge(ws, tl, 1, 7); hcell(ws, tl, 1, '스카치 더블 (남1 + 여1 · 2게임)', grey, '000000', 10); bordered(ws, tl, 2, 7)
    merge(ws, tl, 8, 13); hcell(ws, tl, 8, '베이커 (3인 · 2게임)', grey, '000000', 10); bordered(ws, tl, 9, 13)
    def team_block(c0, size, sample_teams, hdr_r):
        # c0: 팀 열. 이름 c0+1, 성별 c0+2, 핸디 c0+3, 비고 c0+4 ~ c0+5(스카치는 F:G, 베이커는 L:M)
        cN, cG, cH, cB1, cB2 = c0 + 1, c0 + 2, c0 + 3, c0 + 4, c0 + 5
        hcell(ws, hdr_r, c0, '팀'); hcell(ws, hdr_r, cN, '이름'); hcell(ws, hdr_r, cG, '성별'); hcell(ws, hdr_r, cH, '핸디'); merge(ws, hdr_r, cB1, cB2); hcell(ws, hdr_r, cB1, '비고'); hcell(ws, hdr_r, cB2)
        r = hdr_r + 1
        for t in range(TEAMS):
            merge(ws, r, c0, c0, r + size - 1); no = ws.cell(row=r, column=c0, value=t + 1); no.font = f(9, color='808080'); no.alignment = center
            for m in range(size):
                ws.cell(row=r + m, column=c0).border = box
                icell(ws, r + m, cN, None, left, '@'); icell(ws, r + m, cG); icell(ws, r + m, cH); merge(ws, r + m, cB1, cB2); icell(ws, r + m, cB1, None, left, '@'); ws.cell(row=r + m, column=cB2).border = box
                if example and t < len(sample_teams) and m < len(sample_teams[t]):
                    for c, v in zip((cN, cG, cH, cB1), sample_teams[t][m]):
                        if v is not None: ws.cell(row=r + m, column=c, value=v)
            r += size
        dv('"남,여"', f'{L(cG)}{hdr_r + 1}:{L(cG)}{r - 1}')
        return r
    sc_end = team_block(2, 2, [[('홍길동', '남', None, None), ('김영희', '여', 15, None)], [('박철수', '남', 3, None), ('이순자', '여', 20, None)], [('최민수', '남', -7, '프로'), ('정미라', '여', 15, None)]], tl + 1)
    bk_end = team_block(8, 3, [[('홍길동', '남', None, None), ('박철수', '남', None, None), ('최민수', '남', None, None)], [('김영희', '여', None, None), ('이순자', '여', None, None), ('정미라', '여', 5, '여성 3명 +5')]], tl + 1)
    for r in range(tl + 1, bk_end): ws.cell(row=r, column=1).border = box
    memo(ws, sc_end, 2, 7, '팀 핸디는 팀원 핸디 칸의 합계로 적용됩니다. 베이커는 여성 1명 +3, 2명 이상 +5 (한 명 칸에 적어도 됨). 비워 두면 대회에서 규정대로 적용.', bk_end - 1)
    # 하단: 총금액 / 입금계좌 / 신청 문구
    r = bk_end
    merge(ws, r, 1, 2); hcell(ws, r, 1, '총금액', gold, '000000', 10)
    merge(ws, r, 3, 4); tt = ws.cell(row=r, column=3, value=f'=M{tr}'); tt.font = f(10, True); tt.fill = gold; tt.border = box; tt.alignment = right; tt.number_format = '#,##0'; bordered(ws, r, 4, 4)
    merge(ws, r, 5, LASTCOL); acc = ws.cell(row=r, column=5, value='입금계좌 : (대회 공지의 계좌)'); acc.font = f(9); acc.alignment = left
    merge(ws, r + 1, 1, LASTCOL); fin = ws.cell(row=r + 1, column=1, value='위와 같이 제46회 네이버 볼링 동호회 서울 전국대회 참가를 신청합니다.'); fin.font = f(10, True); fin.alignment = center
    end_row = r + 1
    # 인쇄: A4 세로, 폭 맞춤. 1쪽 = 개인전·참가비·3인조, 2쪽 = 스카치·베이커·총금액 (총 인원 행 뒤에서 쪽 나눔)
    from openpyxl.worksheet.pagebreak import Break
    ws.row_breaks.append(Break(id=sr))
    ws.print_area = f'A1:{L(LASTCOL)}{end_row}'
    ws.page_setup.orientation = 'portrait'; ws.page_setup.paperSize = ws.PAPERSIZE_A4; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins.left = ws.page_margins.right = 0.4; ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.print_options.horizontalCentered = True
    ws.protection.sheet = True
    ws.protection.formatCells = False; ws.protection.formatColumns = False; ws.protection.formatRows = False
    ws.protection.selectLockedCells = False

if __name__ == '__main__':
    wb = Workbook(); build(wb.active, False); wb.active.title = '참가신청서'
    build(wb.create_sheet('작성 예시'), True)
    wb.calculation.fullCalcOnLoad = True
    wb.save('docs/signup-template.xlsx'); print('saved docs/signup-template.xlsx')
